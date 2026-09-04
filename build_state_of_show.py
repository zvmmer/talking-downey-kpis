"""Talking Downey — State of the Show deck.

The "read once and you get the show" deck for sponsors. Sits between
OVERALL_REPORT (5 slides, stats only) and SPONSOR_CONTENT_REPORT (21
slides, category deep-dives) — this one tells the LONGITUDINAL story of
the show in 8 punchy slides.

Slides:
    1. Cover — State of the Show
    2. The receipt — cumulative headline stats
    3. Three months, three identities — June/July/August arc
    4. What earned the reach — top all-time posts
    5. The consistent workhorses — cross-month sustained categories
    6. Where the audience shows up — platform breakdown
    7. What we're intentionally not — brand-safety posture
    8. Three directions — where the sponsor could invest next

Uses report_helpers.py (shared with all other decks).

Usage
-----
    ~/talking-downey-kpis/venv/bin/python build_state_of_show.py
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from report_helpers import (
    PALETTE, LAYOUT_W, LAYOUT_H,
    add_rect, add_text, fmt_int,
    load_all_snapshots, posts_in_window, group_by_category,
    slide_new, slide_header, slide_footer_takeaway,
)

JUNE_START   = "2026-06-08"
JUNE_END     = "2026-06-30"
JULY_START   = "2026-07-01"
JULY_END     = "2026-07-31"
AUGUST_START = "2026-08-01"
AUGUST_END   = "2026-08-31"
TODAY = date.today().isoformat()


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, LAYOUT_W, LAYOUT_H, PALETTE["ink"])
    add_rect(slide, 0, 0, 0.45, LAYOUT_H, PALETTE["sky"])

    add_text(slide, 1.0, 1.2, 11, 0.5,
             "STATE OF THE SHOW",
             size=14, color=PALETTE["sky"], bold=True)
    add_text(slide, 1.0, 1.9, 11, 1.3,
             "Talking Downey", size=64, color=PALETTE["white"],
             bold=True, font="Georgia")
    add_text(slide, 1.0, 3.4, 11, 0.6,
             "12 weeks. 1 podcast episode. 4 platforms. 400+ pieces of proof.",
             size=22, color=PALETTE["cream"], font="Georgia")

    add_text(slide, 1.0, 5.0, 8, 0.4,
             f"Since Jun 8, 2026 · through {TODAY}",
             size=13, color=PALETTE["muted"])
    add_text(slide, 1.0, 5.4, 8, 0.4,
             "The read-once-and-you-get-the-show deck.",
             size=13, color=PALETTE["muted"])

    add_text(slide, 1.0, 6.6, 11, 0.4,
             "Prepared by Zhamir Pascual — Z",
             size=12, color=PALETTE["muted"])


def slide_receipt(prs, all_posts):
    """The 'here's what 12 weeks bought you' headline slide."""
    slide = slide_new(prs)
    slide_header(slide, "The receipt",
                 "Everything the show has shipped since launch",
                 "Cumulative totals from Jun 8 through today. Whatever the individual month, this is the body of work.")

    total_views = sum(int(p.get("view_count") or 0) for p in all_posts)
    total_likes = sum(int(p.get("like_count") or 0) for p in all_posts)
    total_comments = sum(int(p.get("comment_count") or 0) for p in all_posts)
    total_posts = len(all_posts)
    avg_v = total_views // max(1, total_posts)

    # Hero number: total views, centered
    add_text(slide, 0.8, 2.2, 11.7, 1.5,
             fmt_int(total_views), size=110, color=PALETTE["sky"],
             bold=True, font="Georgia", align="center")
    add_text(slide, 0.8, 3.8, 11.7, 0.4,
             "views across every tracked post",
             size=16, color=PALETTE["navy"], align="center")

    # Supporting metrics row
    metrics = [
        (fmt_int(total_posts),   "posts published",  PALETTE["navy"]),
        (fmt_int(total_likes),   "total likes",      PALETTE["cherry"]),
        (fmt_int(total_comments),"total comments",   PALETTE["green"]),
        (fmt_int(avg_v),         "avg views / post", PALETTE["gold"]),
    ]
    card_w, gap = 2.85, 0.2
    total_w = card_w * 4 + gap * 3
    x_start = (LAYOUT_W - total_w) / 2
    y = 4.9

    for i, (val, label, color) in enumerate(metrics):
        x = x_start + i * (card_w + gap)
        add_rect(slide, x, y, card_w, 1.4, PALETTE["white"])
        add_rect(slide, x, y, 0.06, 1.4, color)
        add_text(slide, x + 0.2, y + 0.15, card_w - 0.4, 0.7,
                 val, size=28, color=color, bold=True, font="Georgia")
        add_text(slide, x + 0.2, y + 0.9, card_w - 0.4, 0.3,
                 label.upper(), size=9, color=PALETTE["muted"], bold=True)


def slide_arc(prs, jun_posts, jul_posts, aug_posts):
    """Three months, three identities — the longitudinal arc."""
    slide = slide_new(prs)
    slide_header(slide, "Three months, three identities",
                 "The show's own arc",
                 "The same voice, three different chapters. Each month tried something the last month didn't.")

    def _stats(posts):
        v = sum(int(p.get("view_count") or 0) for p in posts)
        return {"posts": len(posts), "views": v, "avg": v // max(1, len(posts))}

    def _top_cat(posts):
        buckets = group_by_category(posts)
        totals = {c: sum(int(p.get("view_count") or 0) for p in ps)
                  for c, ps in buckets.items() if ps and c != "other"}
        if not totals:
            return "—"
        key = max(totals, key=lambda c: totals[c])
        return key.replace("_", " ").title()

    months = [
        ("JUNE",   "Proof of concept",  "Controversy engine at scale — scandal + council chaos led the reach.",
         jun_posts, PALETTE["cherry"]),
        ("JULY",   "Diversification",    "Military service and the FIFA watch party joined the top of the reach ladder — new lanes, same voice.",
         jul_posts, PALETTE["sky"]),
        ("AUGUST", "Exploration",       "No scandal chased, no viral event. Volume held, reach softened — the show explored formats without a defined pillar.",
         aug_posts, PALETTE["gold"]),
    ]

    card_w = 4.2
    gap = 0.2
    total_w = card_w * 3 + gap * 2
    x_start = (LAYOUT_W - total_w) / 2
    y = 2.3

    for i, (label, tagline, note, posts, color) in enumerate(months):
        x = x_start + i * (card_w + gap)
        s = _stats(posts)
        add_rect(slide, x, y, card_w, 4.4, PALETTE["white"])
        add_rect(slide, x, y, 0.08, 4.4, color)

        add_text(slide, x + 0.25, y + 0.2, card_w - 0.4, 0.4,
                 label, size=13, color=color, bold=True)
        add_text(slide, x + 0.25, y + 0.55, card_w - 0.4, 0.6,
                 tagline, size=22, color=PALETTE["ink"], bold=True, font="Georgia")

        # Stat block
        add_text(slide, x + 0.25, y + 1.4, card_w - 0.4, 0.4,
                 fmt_int(s["views"]) + " views",
                 size=20, color=PALETTE["navy"], bold=True, font="Georgia")
        add_text(slide, x + 0.25, y + 1.85, card_w - 0.4, 0.3,
                 f"{s['posts']} posts · {fmt_int(s['avg'])} avg / post",
                 size=11, color=PALETTE["muted"])

        # Top category
        add_rect(slide, x + 0.25, y + 2.3, card_w - 0.5, 0.02, PALETTE["rule"])
        add_text(slide, x + 0.25, y + 2.4, card_w - 0.4, 0.3,
                 "TOP CATEGORY", size=9, color=PALETTE["muted"], bold=True)
        add_text(slide, x + 0.25, y + 2.7, card_w - 0.4, 0.4,
                 _top_cat(posts), size=14, color=PALETTE["ink"], bold=True)

        # Narrative note
        add_text(slide, x + 0.25, y + 3.3, card_w - 0.4, 1.0,
                 note, size=10, color=PALETTE["ink"])

    slide_footer_takeaway(slide,
        "One episode has spawned three months of distinct content identities. That's what a real editorial engine looks like.")


def slide_top_reach(prs, all_posts):
    """What earned the reach — top posts all-time."""
    slide = slide_new(prs)
    slide_header(slide, "What earned the reach",
                 "The top 8 posts since launch",
                 "Ranked by views across all platforms. These are the pieces that pulled outside the average — the ones that traveled.")

    posts = sorted(all_posts, key=lambda p: -int(p.get("view_count") or 0))[:8]

    col_widths = [0.5, 1.1, 6.4, 1.5, 1.3, 1.4]
    headers = ["#", "PLATFORM", "POST", "VIEWS", "LIKES", "CATEGORY"]
    row_y = 2.4
    row_h = 0.5
    x_base = 0.6

    add_rect(slide, x_base, row_y, sum(col_widths), row_h, PALETTE["navy"])
    cx = x_base
    for hd, w in zip(headers, col_widths):
        add_text(slide, cx + 0.1, row_y + 0.12, w - 0.2, 0.3,
                 hd, size=10, color=PALETTE["cream"], bold=True)
        cx += w

    for i, p in enumerate(posts, 1):
        ry = row_y + row_h + (i - 1) * row_h
        if i % 2 == 0:
            add_rect(slide, x_base, ry, sum(col_widths), row_h, PALETTE["soft"])
        if i <= 3:
            add_rect(slide, x_base, ry, 0.06, row_h, PALETTE["cherry"])

        title = (p.get("title") or p.get("description_excerpt") or "")[:75]
        plat = {"instagram": "IG", "tiktok": "TT", "facebook": "FB", "youtube": "YT"}.get(p.get("platform"), "?")
        cat = p.get("_category", "other").replace("_", " ").title()

        cx = x_base
        add_text(slide, cx + 0.1, ry + 0.12, col_widths[0] - 0.2, 0.3,
                 str(i), size=12, color=PALETTE["cherry"], bold=True)
        cx += col_widths[0]
        add_text(slide, cx + 0.1, ry + 0.12, col_widths[1] - 0.2, 0.3,
                 plat, size=11, color=PALETTE["navy"], bold=True)
        cx += col_widths[1]
        add_text(slide, cx + 0.1, ry + 0.12, col_widths[2] - 0.2, 0.3,
                 title, size=10, color=PALETTE["ink"])
        cx += col_widths[2]
        add_text(slide, cx + 0.1, ry + 0.12, col_widths[3] - 0.2, 0.3,
                 fmt_int(p.get("view_count")), size=12, color=PALETTE["navy"], bold=True)
        cx += col_widths[3]
        add_text(slide, cx + 0.1, ry + 0.12, col_widths[4] - 0.2, 0.3,
                 fmt_int(p.get("like_count")), size=11, color=PALETTE["ink"])
        cx += col_widths[4]
        add_text(slide, cx + 0.1, ry + 0.12, col_widths[5] - 0.2, 0.3,
                 cat[:22], size=9, color=PALETTE["muted"])

    slide_footer_takeaway(slide,
        "The pieces that traveled cut across categories — a scandal, a local retail opening, a military explainer, a mall event. Range is the moat.")


def slide_workhorses(prs, jun_posts, jul_posts, aug_posts):
    """The consistent workhorses — categories that showed up every month."""
    slide = slide_new(prs)
    slide_header(slide, "The workhorses",
                 "What consistently works across every month",
                 "Categories that appeared in June, July, and August. Ranked by cumulative reach. These are the reliable lanes.")

    def _cat_stats(posts, cat):
        rel = [p for p in posts if p.get("_category") == cat]
        v = sum(int(p.get("view_count") or 0) for p in rel)
        return {"n": len(rel), "v": v}

    # Find categories present in all three months
    all_posts_by_month = [jun_posts, jul_posts, aug_posts]
    cats_by_month = [set(p.get("_category") for p in m if p.get("_category")) for m in all_posts_by_month]
    persistent = cats_by_month[0] & cats_by_month[1] & cats_by_month[2] - {"other"}

    scored = []
    for cat in persistent:
        totals = [_cat_stats(m, cat) for m in all_posts_by_month]
        v_total = sum(t["v"] for t in totals)
        n_total = sum(t["n"] for t in totals)
        scored.append({
            "cat": cat,
            "n_total": n_total,
            "v_total": v_total,
            "avg": v_total // max(1, n_total),
            "by_month": totals,
        })
    scored.sort(key=lambda s: -s["v_total"])
    top = scored[:6]

    label_map = {
        "food":                 ("Food",                 PALETTE["green"]),
        "community_local":      ("Community Local",      PALETTE["ink"]),
        "community_events":     ("Community Events",     PALETTE["ink"]),
        "downtown_development": ("Downtown Development", PALETTE["gold"]),
        "politics_other":       ("Politics (other)",     PALETTE["navy"]),
        "judicial":             ("Judicial",             PALETTE["navy"]),
        "resident_highlight":   ("Resident Highlight",   PALETTE["green"]),
    }

    card_w, card_h, gap = 3.95, 1.7, 0.2
    x0 = (LAYOUT_W - card_w * 3 - gap * 2) / 2
    y0 = 2.4

    for i, s in enumerate(top):
        col = i % 3
        row = i // 3
        x = x0 + col * (card_w + gap)
        y = y0 + row * (card_h + gap + 0.15)
        label, color = label_map.get(s["cat"], (s["cat"].title(), PALETTE["muted"]))

        add_rect(slide, x, y, card_w, card_h, PALETTE["white"])
        add_rect(slide, x, y, 0.06, card_h, color)
        add_text(slide, x + 0.2, y + 0.15, card_w - 0.4, 0.3,
                 label.upper(), size=10, color=color, bold=True)
        add_text(slide, x + 0.2, y + 0.45, card_w - 0.4, 0.5,
                 fmt_int(s["v_total"]) + " views",
                 size=20, color=PALETTE["navy"], bold=True, font="Georgia")
        add_text(slide, x + 0.2, y + 0.95, card_w - 0.4, 0.3,
                 f"{s['n_total']} posts · {fmt_int(s['avg'])} avg / post",
                 size=10, color=PALETTE["muted"])

        # Per-month strip
        month_labels = ["JUN", "JUL", "AUG"]
        strip_x = x + 0.2
        strip_y = y + 1.3
        strip_w = (card_w - 0.4) / 3
        for j, (ml, t) in enumerate(zip(month_labels, s["by_month"])):
            mx = strip_x + j * strip_w
            add_text(slide, mx, strip_y, strip_w, 0.2,
                     ml, size=8, color=PALETTE["muted"], bold=True)
            add_text(slide, mx, strip_y + 0.2, strip_w, 0.2,
                     f"{t['n']} posts · {fmt_int(t['v'])} views",
                     size=8, color=PALETTE["ink"])

    slide_footer_takeaway(slide,
        "The show's floor is Food + Community Local. Every month, no exception. Everything else is optional flavoring on top of this base.")


def slide_platforms(prs, all_posts):
    """Where the audience shows up."""
    slide = slide_new(prs)
    slide_header(slide, "Where the audience shows up",
                 "Reach by platform, all-time",
                 "Each platform is doing something different for the show — this isn't just cross-posting, it's four distinct audiences.")

    plat_info = [
        ("tiktok",    "TikTok",    PALETTE["ink"],    "The breakout engine — highest per-post reach, home of the viral moments."),
        ("instagram", "Instagram", PALETTE["cherry"], "The workhorse — steady volume, steady reach, the show's Rolodex."),
        ("facebook",  "Facebook",  PALETTE["navy"],   "The reliable floor — held up strongest in August when the others dipped."),
        ("youtube",   "YouTube",   PALETTE["gold"],   "The compounding asset — each episode is a permanent link the algorithm keeps finding."),
    ]

    card_w, card_h, gap = 5.9, 1.6, 0.3
    x0 = (LAYOUT_W - card_w * 2 - gap) / 2
    y0 = 2.4

    for i, (key, label, color, blurb) in enumerate(plat_info):
        col = i % 2
        row = i // 2
        x = x0 + col * (card_w + gap)
        y = y0 + row * (card_h + gap)

        plat_posts = [p for p in all_posts if p.get("platform") == key]
        v = sum(int(p.get("view_count") or 0) for p in plat_posts)
        n = len(plat_posts)
        avg = v // max(1, n)

        add_rect(slide, x, y, card_w, card_h, PALETTE["white"])
        add_rect(slide, x, y, 0.08, card_h, color)
        add_text(slide, x + 0.25, y + 0.15, card_w - 0.4, 0.35,
                 label.upper(), size=12, color=color, bold=True)
        add_text(slide, x + 0.25, y + 0.5, 3.0, 0.6,
                 fmt_int(v), size=30, color=PALETTE["navy"], bold=True, font="Georgia")
        add_text(slide, x + 3.4, y + 0.55, 2.2, 0.3,
                 f"{n} posts", size=13, color=PALETTE["ink"], bold=True)
        add_text(slide, x + 3.4, y + 0.9, 2.4, 0.3,
                 f"{fmt_int(avg)} avg / post", size=11, color=PALETTE["muted"])
        add_text(slide, x + 0.25, y + 1.15, card_w - 0.4, 0.4,
                 blurb, size=10, color=PALETTE["ink"])

    slide_footer_takeaway(slide,
        "Four platforms, four jobs. A sponsor picking Talking Downey isn't buying one audience — they're buying four connected ones.")


def slide_brand_posture(prs, all_posts):
    """What we're intentionally NOT — the brand-safety statement."""
    slide = slide_new(prs)
    slide_header(slide, "Brand posture",
                 "What Talking Downey is intentionally NOT",
                 "A show that runs on scandal has a ceiling. This one has been architected to have a floor instead.")

    scandal_cats = {"trujillo_scandal", "lisette_scandal", "council_chaos",
                    "ai_fake_news", "ice_immigration"}
    total_posts = len(all_posts)
    n_scandal = sum(1 for p in all_posts if p.get("_category") in scandal_cats)
    pct_scandal = (n_scandal / max(1, total_posts)) * 100

    positive_cats = {"food", "community_local", "community_events",
                     "resident_highlight", "downtown_development",
                     "american_pride", "military_service", "fifa_positive",
                     "brand_growth", "political_individual_highlight"}
    n_positive = sum(1 for p in all_posts if p.get("_category") in positive_cats)
    pct_positive = (n_positive / max(1, total_posts)) * 100

    postures = [
        ("NOT a doomscroll show",
         f"{pct_scandal:.0f}% of posts are investigative / scandal-adjacent — the rest are city stories, food, community, and civic life.",
         PALETTE["cherry"]),
        ("NOT a one-note engine",
         f"{pct_positive:.0f}% of the catalog is positive / community-forward content. Sponsors are buying breadth, not controversy.",
         PALETTE["green"]),
        ("NOT reliant on any single platform",
         "Reach is spread across TikTok, Instagram, Facebook, and YouTube. No algorithm change breaks the show.",
         PALETTE["sky"]),
        ("NOT chasing every trend",
         "August proved it — the show can take a quiet month and let the catalog work in the background while it explores.",
         PALETTE["gold"]),
    ]

    y0 = 2.3
    for i, (title, blurb, color) in enumerate(postures):
        y = y0 + i * 1.15
        add_rect(slide, 0.8, y, 11.7, 1.0, PALETTE["white"])
        add_rect(slide, 0.8, y, 0.08, 1.0, color)
        add_text(slide, 1.05, y + 0.12, 11, 0.35,
                 title, size=14, color=color, bold=True, font="Georgia")
        add_text(slide, 1.05, y + 0.5, 11, 0.45,
                 blurb, size=11, color=PALETTE["ink"])

    slide_footer_takeaway(slide,
        "This is what makes the show sponsor-safe. The floor is community. Any heat that shows up is optional flavor.")


def slide_directions(prs, all_posts):
    """Three concrete directions for the sponsor's next move."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, LAYOUT_W, LAYOUT_H, PALETTE["ink"])
    add_rect(slide, 0, 5.0, LAYOUT_W, LAYOUT_H - 5.0, PALETTE["sky"])

    add_text(slide, 0.8, 0.6, 11, 0.4,
             "WHAT COMES NEXT", size=12, color=PALETTE["sky"], bold=True)
    add_text(slide, 0.8, 1.0, 11.5, 1.0,
             "Three directions for the next chapter",
             size=32, color=PALETTE["white"], bold=True, font="Georgia")
    add_text(slide, 0.8, 2.1, 11.5, 0.5,
             "Pick one, pick two, pick all three. Each is grounded in what the last 12 weeks proved.",
             size=13, color=PALETTE["cream"])

    directions = [
        ("LEAN INTO WHAT WORKS",
         "Sponsor a Food + Community Local weekly slot. Sustained per-post reach across all 3 months. Low-risk, high-reliability.",
         PALETTE["green"]),
        ("REOPEN THE ENGINE",
         "Green-light one investigative piece per month. Scandal drove June's breakout — proven ceiling is much higher than the current floor.",
         PALETTE["cherry"]),
        ("TEST A NEW FORMAT",
         "Fund a two-week experiment (YouTube long-form, TikTok series, live stream). The show has shown it can take exploration months without breaking.",
         PALETTE["gold"]),
    ]

    y = 2.9
    for label, blurb, color in directions:
        add_rect(slide, 0.8, y, 11.7, 0.65, PALETTE["cream"])
        add_rect(slide, 0.8, y, 0.08, 0.65, color)
        add_text(slide, 1.05, y + 0.08, 4.5, 0.3,
                 label, size=13, color=color, bold=True)
        add_text(slide, 5.8, y + 0.15, 6.6, 0.4,
                 blurb, size=10, color=PALETTE["ink"])
        y += 0.75

    add_text(slide, 0.8, 5.4, 11.5, 0.5,
             "THE ASK", size=12, color=PALETTE["ink"], bold=True)
    add_text(slide, 0.8, 5.8, 11.5, 1.0,
             "What do you want the next 12 weeks to look like?",
             size=26, color=PALETTE["white"], bold=True, font="Georgia")


# ── main ────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path,
                    default=Path(__file__).parent / "STATE_OF_THE_SHOW.pptx")
    ap.add_argument("--kpis-root", type=Path,
                    default=Path(__file__).parent)
    args = ap.parse_args()

    snaps = load_all_snapshots(args.kpis_root / "episodes_public")
    if not snaps:
        sys.exit(f"No snapshots found under {args.kpis_root/'episodes_public'}")
    print(f"Loaded {len(snaps)} snapshots")

    jun_posts    = posts_in_window(snaps, JUNE_START,   JUNE_END)
    jul_posts    = posts_in_window(snaps, JULY_START,   JULY_END)
    aug_posts    = posts_in_window(snaps, AUGUST_START, AUGUST_END)
    all_posts    = posts_in_window(snaps, JUNE_START,   AUGUST_END)
    print(f"  June:   {len(jun_posts)} posts")
    print(f"  July:   {len(jul_posts)} posts")
    print(f"  August: {len(aug_posts)} posts")
    print(f"  Total:  {len(all_posts)} posts since launch")

    prs = Presentation()
    prs.slide_width  = Inches(LAYOUT_W)
    prs.slide_height = Inches(LAYOUT_H)

    slide_title(prs)
    slide_receipt(prs, all_posts)
    slide_arc(prs, jun_posts, jul_posts, aug_posts)
    slide_top_reach(prs, all_posts)
    slide_workhorses(prs, jun_posts, jul_posts, aug_posts)
    slide_platforms(prs, all_posts)
    slide_brand_posture(prs, all_posts)
    slide_directions(prs, all_posts)

    prs.save(args.out)
    print(f"✓ Wrote: {args.out}")
    print(f"  {len(prs.slides)} slides")
    print(f"  Open with: open '{args.out}'")


if __name__ == "__main__":
    main()
